import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Sana", "Turi", "Kategoriya", "Summa", "Valyuta", "Izoh"]


def build_excel(records: list[dict], chat_id: int) -> str:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Hisobot"

    worksheet.append(HEADERS)
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    for col_idx in range(1, len(HEADERS) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total_income = 0.0
    total_expense = 0.0
    for record in records:
        worksheet.append([record.get(h, "") for h in HEADERS])
        try:
            amount = float(record.get("Summa", 0) or 0)
        except (TypeError, ValueError):
            amount = 0.0
        record_type = str(record.get("Turi", "")).strip().lower()
        if record_type == "kirim":
            total_income += amount
        elif record_type == "chiqim":
            total_expense += amount

    worksheet.append([])
    worksheet.append(["", "", "Jami kirim:", total_income])
    worksheet.append(["", "", "Jami chiqim:", total_expense])
    worksheet.append(["", "", "Sof qoldiq:", total_income - total_expense])

    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = max([len(header)] + [len(str(r.get(header, ""))) for r in records])
        worksheet.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    file_path = Path(tempfile.gettempdir()) / f"hisobot_{chat_id}.xlsx"
    workbook.save(file_path)
    return str(file_path)
